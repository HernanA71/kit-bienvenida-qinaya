# -*- coding: utf-8 -*-
"""
Script para generar el documento oficial en Word (.docx) del:
'Módulo de Computación en la Nube: Capacitación Docente y Revitalización con Sentido Pedagógico'
Para Qinaya.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Establece el color de fondo de una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Establece márgenes internos de una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout_box(doc, title, text, border_color="1E5BA8", bg_color="F1F5F9"):
    """Crea una caja destacada tipo callout con borde lateral izquierdo."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Borde izquierdo grueso
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(title)
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor.from_string(border_color)
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    run_b = p2.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(51, 65, 85)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(4)

def build_docx(filename="Modulo_de_Computacion_en_la_Nube_Qinaya.docx"):
    doc = docx.Document()
    
    # Configuración de márgenes
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
        
        # Header y Footer
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("QINAYA | MÓDULO DE COMPUTACIÓN EN LA NUBE")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Capacitación Docente, Empleabilidad y Revitalización con Propósito • 2026")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # PORTADA / ENCABEZADO PRINCIPAL
    # ==========================================
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(10)
    p_pre.paragraph_format.space_after = Pt(4)
    run_badge = p_pre.add_run("QINAYA EDUCACIÓN DIGITAL • INNOVACIÓN CURRICULAR")
    run_badge.font.name = "Calibri"
    run_badge.font.size = Pt(10)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(234, 88, 12) # Qinaya Orange

    # Título Principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("Módulo de Computación en la Nube")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 91, 168) # Primary Blue

    # Subtítulo
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("De la Caja al Centro de Datos: Revitalización Escolar, Empleabilidad Juvenil y Ciudadanía Digital")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Tabla Ficha Metadatos
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Audiencia y Destinatarios:", "Docentes de Tecnología, Informática y Líderes de Clubes Técnicos Escolares"),
        ("Duración e Intensidad:", "3 Sesiones Formativas (7.5 Horas de Taller Práctico y Pedagógico)"),
        ("Filosofía de Aprendizaje:", "Abrir con el Mundo (IA, Streaming, Realidad) y Cerrar con el Oficio (Empleabilidad TI)"),
        ("Ecosistema Tecnológico:", "Qinaya Linux (~450MB RAM en reposo), Computador Virtual Seguro y Almacenamiento Institucional")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "F1F5F9")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(30, 91, 168)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ==========================================
    # 1. INTRODUCCIÓN Y ENFOQUE
    # ==========================================
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    r_h1 = h1.add_run("1. El Diagnóstico y el Enfoque Pedagógico: Más Allá del Manual")
    r_h1.font.name = "Calibri"
    r_h1.font.color.rgb = RGBColor(30, 91, 168)

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(10)
    p_intro.add_run(
        "Enseñar tecnología no puede reducirse a memorizar una secuencia de botones (abrir Rufus, presionar F9 o sincronizar carpetas). "
        "Cuando un estudiante de noveno o décimo grado entra a la sala de informática, necesita una historia con sentido: ¿Por qué esto le importa a su vida hoy? "
        "¿Cómo se conecta con la música que escucha, los videojuegos que juega o la Inteligencia Artificial de la que todos hablan? Y, sobre todo, "
        "¿cómo le abre las puertas a su futuro laboral?\n\n"
        "Este módulo está diseñado bajo una premisa pedagógica clara: CADA SESIÓN ABRE CON EL MUNDO Y CIERRA CON EL OFICIO. "
        "Primero, se contextualiza por qué la computación dejó de estar atrapada 'dentro de la caja' física del PC y cómo la nube domina el panorama actual. "
        "Luego, se aterriza en el oficio técnico real: el soporte de infraestructura, la administración básica de servicios en la nube y la ciberseguridad "
        "son hoy los escalones de entrada más accesibles y demandados para el primer empleo de un joven graduado. El paso a paso técnico queda, "
        "pero como un medio para empoderar al estudiante, no como un fin mecánico."
    )

    add_callout_box(
        doc,
        "Articulación con los Clubes Técnicos Escolares:",
        "Los colegios cuentan con docentes y Clubes Técnicos que ya enseñan a desarmar, limpiar y reparar hardware. "
        "Este módulo no compite con ese conocimiento: arranca exactamente donde el Club Técnico termina. "
        "El piloto de Qinaya con la Secretaría de Educación de Bogotá (SED) dejó una lección contundente: 'La nube no salva un equipo con hardware en mal estado'. "
        "Si un disco tiene sectores dañados o la placa madre falla, la nube no hace milagros. "
        "Por eso, el triaje técnico se convierte en un ejercicio de pensamiento crítico e ingeniería: decidir con rigor si un computador es candidato viable "
        "para una segunda vida o si debe canalizarse a reciclaje responsable de e-waste.",
        border_color="0284C7",
        bg_color="E0F2FE"
    )

    p_librada = doc.add_paragraph()
    p_librada.paragraph_format.space_after = Pt(12)
    p_librada.add_run(
        "Como prueba de concepto real, el módulo adopta casos tangibles de colegios públicos de Bogotá. El ejemplo del Colegio Santa Librada, "
        "donde un aula de 36 computadores revitalizados lleva más de 24 meses de servicio ininterrumpido sin fallas críticas, "
        "sirve como testimonio vivo para los estudiantes de que la economía circular y la computación en la nube transforman realidades."
    )

    # Parámetros técnicos unificados
    add_callout_box(
        doc,
        "Estándar Técnico Unificado para Revitalización Qinaya:",
        "• Requisitos mínimos de la máquina física candidata: Procesador de 2 núcleos (arquitectura x64), mínimo 1 GB de memoria RAM (óptimo 2 GB) y al menos 10 GB a 16 GB de almacenamiento en disco.\n"
        "• Consumo del sistema base: Qinaya Linux requiere apenas ~450 MB de memoria RAM en reposo, garantizando arranques en menos de 35 a 40 segundos en discos mecánicos convencionales.\n"
        "• Seguridad operativa: Incompatibilidad nativa con malware y ejecutables típicos de Windows (.exe, .vbs, gusanos de USB), protegiendo la sala de informática de forma preventiva.",
        border_color="EA580C",
        bg_color="FFF7ED"
    )

    # ==========================================
    # 2. SESIÓN 1
    # ==========================================
    doc.add_page_break()
    h2_s1 = doc.add_heading(level=1)
    h2_s1.paragraph_format.space_before = Pt(14)
    h2_s1.paragraph_format.space_after = Pt(6)
    r_s1 = h2_s1.add_run("2. Sesión 1: ¿Dónde Vive Hoy la Computación? (Nube, IA y de la Caja al Centro de Datos)")
    r_s1.font.name = "Calibri"
    r_s1.font.color.rgb = RGBColor(30, 91, 168)

    p_s1_meta = doc.add_paragraph()
    p_s1_meta.paragraph_format.space_after = Pt(10)
    p_s1_meta.add_run("Duración: ").bold = True
    p_s1_meta.add_run("2 Horas  |  ")
    p_s1_meta.add_run("El Gancho: ").bold = True
    p_s1_meta.add_run("El estudiante descubre que ya usa la nube todos los días sin saberlo. Del celular al supercomputador remoto.")

    # 2.1 Objetivos de Aprendizaje
    h3 = doc.add_heading(level=2)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(6)
    r_obj1 = h3.add_run("2.1 Objetivos de Aprendizaje")
    r_obj1.font.color.rgb = RGBColor(2, 132, 199)

    obj_s1_table = doc.add_table(rows=2, cols=2)
    obj_s1_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_c0, h_c1 = obj_s1_table.rows[0].cells[0], obj_s1_table.rows[0].cells[1]
    h_c0.width, h_c1.width = Inches(3.25), Inches(3.25)
    set_cell_background(h_c0, "1E5BA8")
    set_cell_background(h_c1, "EA580C")
    set_cell_margins(h_c0, 80, 80, 100, 100)
    set_cell_margins(h_c1, 80, 80, 100, 100)
    h_c0.paragraphs[0].add_run("Objetivos para el Docente").font.bold = True
    h_c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    h_c1.paragraphs[0].add_run("Objetivos de Transferencia a Estudiantes").font.bold = True
    h_c1.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    b_c0, b_c1 = obj_s1_table.rows[1].cells[0], obj_s1_table.rows[1].cells[1]
    set_cell_background(b_c0, "F8FAFC")
    set_cell_background(b_c1, "FFF7ED")
    set_cell_margins(b_c0, 100, 100, 120, 120)
    set_cell_margins(b_c1, 100, 100, 120, 120)
    
    b_c0.paragraphs[0].add_run(
        "• Guiar la discusión reflexiva sobre cómo la nube y la IA redefinieron el concepto de hardware.\n"
        "• Conectar el caso del Colegio Santa Librada como evidencia de sostenibilidad y éxito en el sector público.\n"
        "• Explicar la arquitectura dual de Qinaya Linux: soporte local para robótica escolar y acceso directo a la nube."
    ).font.size = Pt(9.5)
    
    b_c1.paragraphs[0].add_run(
        "• Descubrir que sus aplicaciones favoritas (TikTok, Spotify, ChatGPT) no corren en su teléfono, sino en centros de datos remotos.\n"
        "• Entender qué es un Computador Virtual y por qué no depende de la 'caja' vieja que tienen al frente.\n"
        "• Reconocer el rol del Técnico de Soporte Cloud como una opción real y alcanzable de empleo formal."
    ).font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 2.2 Desarrollo Temático
    h3 = doc.add_heading(level=2)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(6)
    r_conc1 = h3.add_run("2.2 Desarrollo Temático: La Descentralización del Cómputo")
    r_conc1.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph(
        "A. El Mito de 'La Caja': ¿Dónde Ocurre la Magia?\n"
        "Durante décadas se enseñó que el poder de un computador dependía del tamaño de su torre y la marca de sus chips. "
        "Hoy, la revolución de los modelos de Inteligencia Artificial y las plataformas de streaming ha cambiado las reglas del juego: "
        "los cálculos pesados ocurren en centros de datos distribuidos en todo el planeta. La máquina que el usuario toca es simplemente una "
        "interfaz de visualización y control (Cliente Ligero o Thin Client).\n\n"
        "B. El Caso del Colegio Santa Librada (Bogotá):\n"
        "En lugar de hablar de teorías abstractas sobre sistemas operativos, el docente presenta la historia real de la sala de cómputo del Colegio "
        "Santa Librada: 36 equipos descartados que fueron reactivados con la solución de Qinaya y que hoy acumulan más de dos años de trabajo "
        "escolar continuo, ejecutando software educativo sin que los colegios hayan gastado en renovar sus torres.\n\n"
        "C. Qinaya Linux y las Capacidades Locales:\n"
        "Un sistema operativo liviano de última generación no limita al estudiante: le da autonomía. Qinaya Linux mantiene un entorno local con navegador web "
        "ligero y soporte nativo para puertos USB, permitiendo programar placas físicas de robótica escolar como Arduino o BBC Micro:bit, "
        "mientras que su lanzador institucional (Launcher) abre la puerta a un computador virtual con toda la suite de ofimática y software educativo.\n\n"
        "D. El Oficio: El Primer Escalón en Tecnología:\n"
        "El soporte de computación en la nube (Cloud Help Desk) es una de las áreas con mayor escasez de talento junior en América Latina. "
        "Aprender cómo se autentica una sesión remota, cómo se gestiona la latencia y cómo se conectan los periféricos es la antesala técnica "
        "para certificaciones internacionales y empleo formal."
    )

    # 2.3 Reto Estudiantil #1
    add_callout_box(
        doc,
        "Reto Estudiantil #1: 'El Mapa de la Computación Oculta'",
        "Actividad de Aula para Estudiantes de 9º/10º:\n"
        "1. Reflexión en parejas: Mapear 3 aplicaciones que usan a diario (ej. Instagram, Free Fire, Gemini) y responder: ¿Qué parte se ejecuta en mi pantalla y qué parte viaja a un servidor en la nube?\n"
        "2. Encendido del PC con Qinaya Linux y cronometraje de arranque (verificar arranque veloz en menos de 40 segundos).\n"
        "3. Conectar una placa de robótica (Arduino/Micro:bit) o diadema y verificar en el sistema local su reconocimiento inmediato.\n"
        "4. Ejecutar el Launcher institucional de Qinaya e ingresar a su Computador Virtual.\n"
        "5. Comparar la respuesta del cursor y la resolución frente a una máquina local pesada.\n\n"
        "Entregable del Estudiante: Infografía o esquema 'De mi Aula al Centro de Datos' + Registro de verificación de acceso a su PC Virtual.",
        border_color="EA580C",
        bg_color="FFF7ED"
    )

    # ==========================================
    # 3. SESIÓN 2
    # ==========================================
    doc.add_page_break()
    h2_s2 = doc.add_heading(level=1)
    h2_s2.paragraph_format.space_before = Pt(14)
    h2_s2.paragraph_format.space_after = Pt(6)
    r_s2 = h2_s2.add_run("3. Sesión 2: Misión Rescate (Economía Circular, Triaje Crítico e Instalación)")
    r_s2.font.name = "Calibri"
    r_s2.font.color.rgb = RGBColor(30, 91, 168)

    p_s2_meta = doc.add_paragraph()
    p_s2_meta.paragraph_format.space_after = Pt(10)
    p_s2_meta.add_run("Duración: ").bold = True
    p_s2_meta.add_run("3 Horas  |  ")
    p_s2_meta.add_run("El Enfoque: ").bold = True
    p_s2_meta.add_run("Pensamiento crítico técnico: la nube no salva hardware roto. Economía circular y puesta a punto con Rufus.")

    # 3.1 Objetivos de Aprendizaje
    h3 = doc.add_heading(level=2)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(6)
    r_obj2 = h3.add_run("3.1 Objetivos de Aprendizaje")
    r_obj2.font.color.rgb = RGBColor(2, 132, 199)

    obj_s2_table = doc.add_table(rows=2, cols=2)
    obj_s2_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_c0, h_c1 = obj_s2_table.rows[0].cells[0], obj_s2_table.rows[0].cells[1]
    h_c0.width, h_c1.width = Inches(3.25), Inches(3.25)
    set_cell_background(h_c0, "1E5BA8")
    set_cell_background(h_c1, "EA580C")
    set_cell_margins(h_c0, 80, 80, 100, 100)
    set_cell_margins(h_c1, 80, 80, 100, 100)
    h_c0.paragraphs[0].add_run("Objetivos para el Docente").font.bold = True
    h_c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    h_c1.paragraphs[0].add_run("Objetivos de Transferencia a Estudiantes").font.bold = True
    h_c1.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    b_c0, b_c1 = obj_s2_table.rows[1].cells[0], obj_s2_table.rows[1].cells[1]
    set_cell_background(b_c0, "F8FAFC")
    set_cell_background(b_c1, "FFF7ED")
    set_cell_margins(b_c0, 100, 100, 120, 120)
    set_cell_margins(b_c1, 100, 100, 120, 120)
    
    b_c0.paragraphs[0].add_run(
        "• Articular las habilidades de diagnóstico de hardware previas del Club Técnico con la viabilidad para la nube.\n"
        "• Enseñar el criterio de decisión: distinguir entre un equipo recuperable (2 núcleos, 1GB RAM) y chatarra no viable.\n"
        "• Dominar el proceso de flasheo masivo con Rufus e instalación limpia de Qinaya Linux."
    ).font.size = Pt(9.5)
    
    b_c1.paragraphs[0].add_run(
        "• Comprender el impacto ambiental de la basura electrónica (e-waste) y la importancia de la economía circular.\n"
        "• Aprender a tomar decisiones técnicas reales: emitir un veredicto de viabilidad sobre una máquina física.\n"
        "• Adquirir la destreza práctica de configurar BIOS/Boot Menu e instalar un sistema operativo libre."
    ).font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3.2 Desarrollo Temático
    h3 = doc.add_heading(level=2)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(6)
    r_conc2 = h3.add_run("3.2 Desarrollo Temático: Criterio Técnico y Despliegue en Campo")
    r_conc2.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph(
        "A. La Gran Lección del Piloto SED: La Nube no Salva Hardware Roto\n"
        "Uno de los hallazgos más valiosos del trabajo de campo de Qinaya en instituciones distritales de Bogotá fue desmentir la creencia de que "
        "'la nube revive cualquier cosa'. Si un disco duro mecánico presenta sectores dañados (bad sectors) que generan bloqueos de lectura, "
        "si los módulos de memoria RAM arrojan errores de paridad o si el procesador se apaga por recalentamiento de pasta térmica vencida, "
        "el software en la nube se congelará. La nube necesita una base física sana.\n\n"
        "B. El Triaje Crítico: Decisión de Ingeniería\n"
        "Por eso, antes de instalar, el estudiante aplica un triaje riguroso. Para que un computador sea candidato a una segunda vida debe cumplir:\n"
        "• Procesador: Mínimo 2 núcleos con soporte de instrucciones de 64 bits.\n"
        "• Memoria RAM: Al menos 1 GB físico funcional (2 GB recomendado para multitarea local fluida).\n"
        "• Disco de almacenamiento: Al menos 10 GB a 16 GB en estado de salud S.M.A.R.T. aceptable.\n"
        "Si el equipo no supera este umbral, el criterio profesional del estudiante dicta que no debe forzarse: se aprovechan sus piezas útiles "
        "para reparar otras máquinas y el resto se deriva a reciclaje ambiental certificado.\n\n"
        "C. El Procedimiento Técnico con Rufus (El Medio para el Fin):\n"
        "Una vez validado el equipo, los estudiantes preparan su medio de rescate con Rufus:\n"
        "1. Seleccionar la memoria USB y cargar la imagen ISO oficial de Qinaya Linux.\n"
        "2. Identificar el esquema de partición: MBR para computadores escolares con BIOS Legacy tradicional, o GPT si la placa soporta UEFI moderno.\n"
        "3. Grabación en modo imagen ISO y verificación.\n"
        "4. Configuración de arranque: Desactivar 'Secure Boot' en el BIOS y acceder al Boot Menu (HP: F9/Esc; Lenovo: F12; Dell: F12; Clones: F8/F12).\n"
        "5. Instalación guiada: Particionado automático Ext4 y creación de partición Swap de soporte de memoria.\n\n"
        "D. El Oficio: Técnico de Despliegue y Mantenimiento de Campo\n"
        "Este ejercicio reproduce exactamente lo que hace un técnico de campo o analista de soporte en una empresa al aprovisionar 50 estaciones de trabajo. "
        "El estudiante no solo instaló un sistema: aprendió despliegue masivo y resolución de incidencias en caliente."
    )

    # 3.3 Reto Estudiantil #2
    add_callout_box(
        doc,
        "Reto Estudiantil #2: 'Operación Rescate: Del Triaje a la Instalación'",
        "Actividad de Aula para Estudiantes de 9º/10º:\n"
        "1. Formar brigadas de soporte técnico de 2 estudiantes y recibir un computador escolar descartado o lento.\n"
        "2. Aplicar el Triaje Crítico: Encender, entrar al BIOS, verificar 2 núcleos, memoria RAM (>1GB) y comprobar ausencia de ruidos mecánicos de disco.\n"
        "3. Emitir el veredicto en el acta técnica: ¿Es candidato viable para revitalización con Qinaya o se canaliza a donación de repuestos?\n"
        "4. En equipos declarados viables: Grabar la memoria USB con Rufus y ejecutar la instalación completa de Qinaya Linux.\n"
        "5. Reiniciar, cronometrar arranque (debe ser menor a 40 segundos) y validar periféricos (red y sonido).\n\n"
        "Entregable del Estudiante: Acta de Rescate Técnico diligenciada con justificación del veredicto, especificaciones finales y firma del equipo técnico.",
        border_color="EA580C",
        bg_color="FFF7ED"
    )

    # ==========================================
    # 4. SESIÓN 3
    # ==========================================
    doc.add_page_break()
    h2_s3 = doc.add_heading(level=1)
    h2_s3.paragraph_format.space_before = Pt(14)
    h2_s3.paragraph_format.space_after = Pt(6)
    r_s3 = h2_s3.add_run("4. Sesión 3: Trabajar en la Nube con Criterio (Colaboración, Ciudadanía Digital y Portafolio)")
    r_s3.font.name = "Calibri"
    r_s3.font.color.rgb = RGBColor(30, 91, 168)

    p_s3_meta = doc.add_paragraph()
    p_s3_meta.paragraph_format.space_after = Pt(10)
    p_s3_meta.add_run("Duración: ").bold = True
    p_s3_meta.add_run("2.5 Horas  |  ")
    p_s3_meta.add_run("El Enfoque: ").bold = True
    p_s3_meta.add_run("Uso ético y profesional de la nube: privacidad de datos personales de menores, colaboración segura y portafolio laboral.")

    # 4.1 Objetivos de Aprendizaje
    h3 = doc.add_heading(level=2)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(6)
    r_obj3 = h3.add_run("4.1 Objetivos de Aprendizaje")
    r_obj3.font.color.rgb = RGBColor(2, 132, 199)

    obj_s3_table = doc.add_table(rows=2, cols=2)
    obj_s3_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_c0, h_c1 = obj_s3_table.rows[0].cells[0], obj_s3_table.rows[0].cells[1]
    h_c0.width, h_c1.width = Inches(3.25), Inches(3.25)
    set_cell_background(h_c0, "1E5BA8")
    set_cell_background(h_c1, "16A34A")
    set_cell_margins(h_c0, 80, 80, 100, 100)
    set_cell_margins(h_c1, 80, 80, 100, 100)
    h_c0.paragraphs[0].add_run("Objetivos para el Docente").font.bold = True
    h_c0.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    h_c1.paragraphs[0].add_run("Objetivos de Transferencia a Estudiantes").font.bold = True
    h_c1.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    b_c0, b_c1 = obj_s3_table.rows[1].cells[0], obj_s3_table.rows[1].cells[1]
    set_cell_background(b_c0, "F8FAFC")
    set_cell_background(b_c1, "F0FDF4")
    set_cell_margins(b_c0, 100, 100, 120, 120)
    set_cell_margins(b_c1, 100, 100, 120, 120)
    
    b_c0.paragraphs[0].add_run(
        "• Fomentar la protección de datos personales de menores conforme a la normativa de privacidad y habeas data escolar.\n"
        "• Posicionar el almacenamiento institucional del Computador Virtual Qinaya como el repositorio protegido principal.\n"
        "• Evaluar y certificar el desempeño del estudiante mediante un portafolio digital tangible."
    ).font.size = Pt(9.5)
    
    b_c1.paragraphs[0].add_run(
        "• Desarrollar hábitos de ciberhigiene y entender por qué no se deben subir datos personales a cualquier plataforma libre.\n"
        "• Trabajar en equipo en tiempo real sobre documentos compartidos dentro del entorno protegido de Qinaya.\n"
        "• Construir su primer Portafolio Técnico Digital para demostrar sus competencias ante empleadores o directivas."
    ).font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 4.2 Desarrollo Temático
    h3 = doc.add_heading(level=2)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(6)
    r_conc3 = h3.add_run("4.2 Desarrollo Temático: Ciudadanía Digital y Gestión Segura en la Nube")
    r_conc3.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph(
        "A. Protección de Datos y Privacidad Escolar (El Cuidado de Menores):\n"
        "En instituciones educativas oficiales, la protección de datos personales de niñas, niños y adolescentes es una obligación ética y legal estricta. "
        "Obligar a los estudiantes a crear cuentas en plataformas externas comerciales o subir tareas a servidores públicos no auditados genera serias "
        "vulnerabilidades de privacidad. Por este motivo, el PROTAGONISTA CENTRAL DE ALMACENAMIENTO es el Computador Virtual de Qinaya, un entorno "
        "institucional cerrado, seguro y blindado donde los trabajos quedan resguardados bajo control del colegio.\n\n"
        "Las herramientas externas (como solicitudes de archivos en Dropbox, repositorios abiertos en MEGA o enlaces efímeros en Wormhole) se abordan "
        "únicamente como conceptos complementarios de transferencia rápida para proyectos públicos, recalcando siempre la regla de oro: "
        "nunca almacenar datos personales, nombres completos, fotos o documentos de identidad en servicios externos gratuitos.\n\n"
        "B. Rigor Técnico: Resiliencia ante el Malware Común:\n"
        "Ningún sistema en el mundo debe publicitarse falsamente como '100% inmune' o 'totalmente libre de virus'. La ventaja técnica real de Qinaya "
        "radica en que la inmensa mayoría del malware que azota las salas escolares de cómputo (gusanos propagados por USBs, scripts de acceso directo, "
        "archivos .vbs y ejecutables .exe maliciosos) fue escrito exclusivamente para sistemas Windows tradicionales. "
        "Al operar sobre el kernel de Qinaya Linux, estos archivos son naturalmente incompatibles e inofensivos, reduciendo drásticamente las fallas "
        "y los tiempos muertos de clase sin ralentizar la máquina con antivirus pesados.\n\n"
        "C. Colaboración en Tiempo Real y Trabajo Productivo:\n"
        "Los estudiantes aprenden a coeditar documentos de trabajo, compartir carpetas institucionales de clase y gestionar versiones sin el caos "
        "de 'documento_final_v2_este_si.docx'.\n\n"
        "D. El Oficio: Portafolio Técnico y Certificación de Competencias:\n"
        "Para un joven que busca sus primeras prácticas laborales, una hoja de vida sin experiencia previa dice poco; un portafolio técnico con "
        "evidencias concretas lo dice todo. La sesión cierra con la consolidación del Portafolio del Estudiante: diagnóstico de hardware, acta de "
        "revitalización y proyectos elaborados en la nube, habilitando la entrega de su Certificado de Aprobación del Módulo."
    )

    # 4.3 Reto Estudiantil #3
    add_callout_box(
        doc,
        "Reto Estudiantil #3: 'El Portafolio del Técnico Cloud Escolar'",
        "Actividad de Aula para Estudiantes de 9º/10º:\n"
        "1. Iniciar sesión en el Computador Virtual Qinaya y acceder a su espacio de almacenamiento institucional protegido.\n"
        "2. En parejas, coeditar en tiempo real un informe técnico sobre 'La Nube como Motor de Sostenibilidad Escolar'.\n"
        "3. Aplicar el protocolo de ciberhigiene: verificar que el documento no contenga datos personales sensibles expuestos.\n"
        "4. Consolidar en una carpeta digital las 3 evidencias del módulo: (a) El mapa de la computación, (b) El acta de triaje e instalación del PC rescatado, y (c) El informe colaborativo.\n"
        "5. Presentar el portafolio para la revisión del docente y la emisión del Certificado de Competencias.\n\n"
        "Entregable del Estudiante: Portafolio Técnico Digital completo y validación de competencias para certificación institucional.",
        border_color="16A34A",
        bg_color="F0FDF4"
    )

    # ==========================================
    # 5. ANEXO FORMATOS
    # ==========================================
    doc.add_page_break()
    h2_anx = doc.add_heading(level=1)
    h2_anx.paragraph_format.space_before = Pt(14)
    h2_anx.paragraph_format.space_after = Pt(6)
    r_anx = h2_anx.add_run("5. Anexo: Formato QN-REV-02 (Ficha de Triaje Crítico y Decisión de Rescate)")
    r_anx.font.name = "Calibri"
    r_anx.font.color.rgb = RGBColor(30, 91, 168)

    doc.add_paragraph(
        "Formato oficial para los Clubes Técnicos y estudiantes en el laboratorio de sistemas. "
        "Enfatiza la toma de decisiones informadas de ingeniería y el pensamiento crítico frente al hardware escolar."
    )

    ficha_table = doc.add_table(rows=8, cols=3)
    ficha_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ficha_headers = ["Punto de Evaluación Técnica", "Especificación Medida en BIOS / Hardware", "Criterio de Viabilidad"]
    for j, htext in enumerate(ficha_headers):
        cell = ficha_table.rows[0].cells[j]
        set_cell_background(cell, "1E5BA8")
        set_cell_margins(cell, 60, 60, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(htext)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    ficha_rows = [
        ("Identificación del Equipo", "N° de Inventario: _________ | Marca: __________________", "Registro Escolar"),
        ("Arquitectura del Procesador", "Modelo: _________________ | Núcleos: [  ] 2 o más   [  ] 1 núcleo", "[  ] Apto (Mín. 2 núcleos x64)\n[  ] No viable"),
        ("Salud de Memoria RAM", "Capacidad: [  ] 1GB   [  ] 2GB   [  ] 4GB+ | Sin errores de arranque", "[  ] Apto (Mín. 1GB / Ópt. 2GB)\n[  ] Memoria dañada"),
        ("Estado Físico del Disco", "Capacidad: _______ GB | Ausencia de ruidos metálicos de cabezal", "[  ] Apto (Mín. 10GB sano)\n[  ] Reemplazar disco"),
        ("VEREDICTO TÉCNICO", "¿El equipo califica para rescate con Qinaya Linux?", "[  ] VIABLE PARA REVITALIZACIÓN\n[  ] DONACIÓN DE REPUESTOS"),
        ("Arranque con Qinaya Linux", "Segundos medidos desde encendido hasta ver escritorio: ______ seg", "[  ] Superado (<40 seg)\n[  ] En revisión"),
        ("Conexión al PC Virtual", "Launcher Qinaya ejecutado y sesión de prueba iniciada con éxito", "[  ] Sesión Estable (<100ms)")
    ]
    for r_idx, r_data in enumerate(ficha_rows):
        row = ficha_table.rows[r_idx + 1]
        for c_idx, c_val in enumerate(r_data):
            cell = row.cells[c_idx]
            bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            if r_idx == 4:
                bg = "FFF7ED" # Destacar veredicto
            set_cell_background(cell, bg)
            set_cell_margins(cell, 50, 50, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(c_val)
            r.font.size = Pt(9)
            if c_idx == 0:
                r.font.bold = True
            if r_idx == 4:
                r.font.bold = True
                if c_idx == 0:
                    r.font.color.rgb = RGBColor(234, 88, 12)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Firmas
    firmas_p = doc.add_paragraph()
    firmas_p.paragraph_format.space_before = Pt(30)
    firmas_p.paragraph_format.space_after = Pt(0)
    firmas_p.add_run(
        "________________________________________                  ________________________________________\n"
        "   Firma del Estudiante Técnico Líder                            Firma del Docente Tutor de Aula"
    ).font.size = Pt(10)

    # Guardar
    doc.save(filename)
    print(f"Documento guardado con éxito en: {filename}")

if __name__ == "__main__":
    build_docx()
